#!/usr/bin/env python3
"""
Cacciatore di occasioni: aste giudiziarie di beni mobili (veicoli, arte/oreficeria/
orologeria/antiquariato) nelle province di Modena, Bologna e Reggio Emilia, andate
deserte >= 2 volte e ancora in vendita sul circuito Astagiudiziaria.com.

Pensato per girare come GitHub Actions (headless, nessuna interazione manuale).

Perche' Playwright e non solo `requests`:
la chiave API pubblica di Typesense usata dal sito viene ruotata periodicamente ed e'
esposta solo lato client durante il caricamento della pagina (non e' garantito trovarla
in modo affidabile nell'HTML statico). Playwright apre davvero la pagina, intercetta la
richiesta di rete verso Typesense e legge il token JWT dal runtime config della pagina,
esattamente come farebbe un browser vero.

Output: un file .xlsx con 3 fogli (Riepilogo, Veicoli, Arte-Oreficeria-Antiquariato),
salvato in ./output/ e caricato su una cartella Google Drive (se le credenziali sono
configurate).
"""
import base64
import io
import json
import os
import re
import sys
from datetime import date, datetime

import requests
from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------
# CONFIGURAZIONE
# --------------------------------------------------------------------------------------

ENTRY_URL = "https://www.ivgmodena.it/"  # basta un solo sito del circuito: il database
                                          # Typesense e' condiviso da tutti gli istituti.

TYPESENSE_HOST = "https://typesense.astagiudiziaria.com"
TYPESENSE_COLLECTION = "astagiudiziaria-prod-v3"
CORE_API_HOST = "https://core-v3.astagiudiziaria.com"

# Mappa provincia/tribunale -> dominio pubblico per raggiungere la scheda del singolo
# annuncio (necessario per leggere cauzione e link ai documenti). Se in futuro cambiano
# gli host, aggiornare qui.
DOMAIN_BY_KEYWORD = [
    ("modena", "https://www.ivgmodena.it"),
    ("bologna", "https://www.ivgbologna.it"),
    ("reggio emilia", "https://ivgreggioemilia.fallcoaste.it"),
]

FILTER_BY = (
    "genre: [MOBILI] && status: [In vendita] && "
    "category: [ARTE -OREFICERIA - OROLOGERIA- ANTIQUARIATO, AUTOVEICOLI E CICLI] && "
    "province: [Bologna, Modena, Reggio Emilia]"
)

MIN_ASTE_DESERTE = 2

OUTPUT_DIR = "output"
REPORT_PREFIX = "Report_Aste_Beni_Mobili_MO_BO_RE"

FONT = "Arial"

BROWSER_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Referer": ENTRY_URL,
    "Origin": ENTRY_URL.rstrip("/"),
}


# --------------------------------------------------------------------------------------
# STEP 1 - Cattura chiave Typesense + token API tramite un browser headless
# --------------------------------------------------------------------------------------

def get_live_tokens():
    """Apre la home page con un browser headless e intercetta:
    - la api-key pubblica di Typesense (dalla query string della richiesta di rete)
    - il JWT 'astaToken' usato per l'API core-v3 (dal runtime config della pagina)
    """
    captured = {}

    def handle_response(response):
        if "typesense.astagiudiziaria.com" in response.url and "x-typesense-api-key=" in response.url and response.status == 200:
            m = re.search(r"x-typesense-api-key=([^&]+)", response.url)
            if m:
                captured["typesense_key"] = m.group(1)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(user_agent=(
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        ))
        page.on("response", handle_response)
        page.goto(ENTRY_URL, wait_until="networkidle", timeout=45000)
        # forza un secondo giro di ricerca se la prima richiesta non e' stata catturata
        if "typesense_key" not in captured:
            page.reload(wait_until="networkidle", timeout=45000)
        try:
            asta_token = page.evaluate("window.__NUXT__.config.astaToken")
        except Exception:
            asta_token = None
        browser.close()

    if "typesense_key" not in captured:
        raise RuntimeError(
            "Non sono riuscito a catturare la api-key di Typesense dalla pagina. "
            "Il sito potrebbe aver cambiato struttura: ispezionare manualmente "
            "le richieste di rete su " + ENTRY_URL
        )
    if not asta_token:
        raise RuntimeError(
            "Non sono riuscito a leggere window.__NUXT__.config.astaToken dalla pagina."
        )

    return captured["typesense_key"], asta_token


# --------------------------------------------------------------------------------------
# STEP 2 - Query Typesense (ricerca lotti) e API storico-inserzione
# --------------------------------------------------------------------------------------

def search_lots(typesense_key):
    url = f"{TYPESENSE_HOST}/collections/{TYPESENSE_COLLECTION}/documents/search"
    params = {
        "query_by": "ivg_short_name,province,numero_procedura,category,subcategory,"
                    "title,inserzioneEspVendita,tags,city",
        "per_page": 250,
        "page": 1,
        "q": "*",
        "filter_by": FILTER_BY,
    }
    headers = {**BROWSER_HEADERS, "X-TYPESENSE-API-KEY": typesense_key}
    r = requests.get(url, params=params, headers=headers, timeout=30)
    r.raise_for_status()
    data = r.json()
    return [h["document"] for h in data.get("hits", [])]


def get_storico(asta_token, lot_id):
    url = f"{CORE_API_HOST}/api/v1/front/inserzioni/storico-inserzione/{lot_id}"
    headers = {**BROWSER_HEADERS, "Authorization": f"Bearer {asta_token}"}
    r = requests.get(url, headers=headers, timeout=20)
    if not r.ok:
        return []
    return r.json().get("data", [])


# --------------------------------------------------------------------------------------
# STEP 3 - Dettagli di pagina (cauzione, documenti) via browser headless
# --------------------------------------------------------------------------------------

def domain_for(text):
    text = (text or "").lower()
    for keyword, domain in DOMAIN_BY_KEYWORD:
        if keyword in text:
            return domain
    return DOMAIN_BY_KEYWORD[0][1]  # fallback: Modena


def scrape_lot_page(page, url):
    """Naviga sulla scheda annuncio e ne estrae cauzione + link a perizia/avviso/ordinanza."""
    try:
        page.goto(url, wait_until="networkidle", timeout=30000)
    except Exception:
        return {"cauzione": None, "links": []}

    result = page.evaluate(
        """
        () => {
          const body = document.body.innerText;
          const m = body.match(/Cauzione[:\\s]*\\u20ac?\\s*([\\d.,]+)/i);
          const links = [...document.querySelectorAll('a')]
            .filter(a => /perizia|avviso|ordinanza/i.test(a.textContent) ||
                         /perizia|avviso|ordinanza/i.test(a.href))
            .map(a => a.href);
          return { cauzione: m ? m[1] : null, links: [...new Set(links)] };
        }
        """
    )
    return result


# --------------------------------------------------------------------------------------
# STEP 4 - Filtro di idoneita'
# --------------------------------------------------------------------------------------

def build_eligible_lots(raw_lots, asta_token):
    eligible = []
    for doc in raw_lots:
        lot_id = doc["id"]
        hist = get_storico(asta_token, lot_id)
        past = [a for a in hist if not a.get("isMe")]
        n_deserte = sum(1 for a in past if a.get("status") == "Non aggiudicato")
        if n_deserte < MIN_ASTE_DESERTE:
            continue
        current = next((a for a in hist if a.get("isMe")), hist[-1] if hist else {})
        first = past[0] if past else {}
        eligible.append({
            "id": lot_id,
            "title": doc.get("title", ""),
            "category": doc.get("category", ""),
            "subcategory": (doc.get("subcategory") or [""])[0] if isinstance(doc.get("subcategory"), list) else doc.get("subcategory"),
            "city": doc.get("city"),
            "province": doc.get("province"),
            "procedura": doc.get("numero_procedura"),
            "tribunale": doc.get("ivg_short_name") or doc.get("tribunal"),
            "n_deserte": n_deserte,
            "n_tentativi": len(hist),
            "prima_asta": first.get("dataOraVendita"),
            "prezzo_prima_asta": first.get("prezzoValoreBase"),
            "prossima_asta": current.get("dataOraVendita"),
            "prezzo_corrente": current.get("prezzoValoreBase"),
            "min_offer": doc.get("minimumOffer"),
            "permalink": current.get("permalink") or doc.get("permalink"),
        })
    return eligible


# --------------------------------------------------------------------------------------
# STEP 5 - Confronto con il report della run precedente (scaricato da Drive)
# --------------------------------------------------------------------------------------

def extract_ids_from_previous_report(xlsx_bytes):
    """Legge la colonna nascosta 'ID' nei fogli Veicoli / Arte-Oreficeria-Antiquariato
    di un report precedente per calcolare new entry / lotti scomparsi."""
    ids = set()
    try:
        wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    except Exception:
        return ids
    for sheet_name in ("Veicoli", "Arte-Oreficeria-Antiquariato"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        if "ID" not in header:
            continue
        id_col = header.index("ID") + 1
        for row in ws.iter_rows(min_row=2):
            val = row[id_col - 1].value
            if val:
                ids.add(str(val))
    return ids


# --------------------------------------------------------------------------------------
# STEP 6 - Costruzione del file Excel
# --------------------------------------------------------------------------------------

COLUMNS = [
    "ID", "Tipologia bene", "Comune", "Provincia", "Procedura/Tribunale",
    "N. aste andate deserte finora", "Data prima asta", "Data prossima asta",
    "Valore di perizia (prima asta)", "Prezzo base asta corrente", "Offerta minima",
    "Cauzione", "Valore stimato da te", "Link scheda annuncio",
    "Link documenti (perizia/avviso/ordinanza)",
]


def eur(v):
    if v in (None, ""):
        return "n.d."
    try:
        return f"€ {float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return str(v)


def build_workbook(eligible, new_ids, removed_ids, previous_found):
    wb = Workbook()
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    normal_font = Font(name=FONT, size=10)
    bold_font = Font(name=FONT, bold=True, size=10)
    title_font = Font(name=FONT, bold=True, size=14)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")

    # ---- Riepilogo ----
    ws = wb.active
    ws.title = "Riepilogo"
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Report Aste Beni Mobili Deserte — Modena, Bologna, Reggio Emilia"
    ws["B2"].font = title_font
    ws["B3"] = f"Data esecuzione: {date.today().strftime('%d/%m/%Y')} (run automatica GitHub Actions)"
    ws["B3"].font = Font(name=FONT, italic=True, size=10)

    veicoli_n = sum(1 for l in eligible if "AUTOVEICOLI" in (l["category"] or ""))
    arte_n = len(eligible) - veicoli_n
    by_prov = {}
    for l in eligible:
        by_prov[l["province"]] = by_prov.get(l["province"], 0) + 1

    rows = [
        ("", ""),
        ("Criterio di idoneita'", f"Lotto in vendita con almeno {MIN_ASTE_DESERTE} aste precedenti concluse 'Non aggiudicato'"),
        ("Lotti idonei totali", len(eligible)),
        ("  di cui Veicoli", veicoli_n),
        ("  di cui Arte-Oreficeria-Orologeria-Antiquariato", arte_n),
        ("  Provincia di Modena", by_prov.get("Modena", 0)),
        ("  Provincia di Bologna", by_prov.get("Bologna", 0)),
        ("  Provincia di Reggio Emilia", by_prov.get("Reggio Emilia", 0)),
        ("", ""),
        ("Confronto con la run precedente", "" if previous_found else "Nessun report precedente trovato su Drive: prima esecuzione utile."),
    ]
    if previous_found:
        rows.append(("  Nuovi lotti (mai visti prima)", len(new_ids)))
        rows.append(("  Lotti scomparsi (probabile aggiudicazione/ritiro)", len(removed_ids)))
        if new_ids:
            rows.append(("  ID nuovi lotti", ", ".join(sorted(new_ids))))
        if removed_ids:
            rows.append(("  ID lotti scomparsi", ", ".join(sorted(removed_ids))))
    rows.append(("", ""))
    rows.append((
        "Limiti della stima",
        "La colonna 'Valore stimato da te' non viene compilata automaticamente da questo "
        "script: richiede una valutazione qualitativa (marca/modello/anno per i veicoli, "
        "tipologia/materiale per arte e oreficeria) che questa routine non esegue. "
        "Compilarla a mano o integrare un modello di stima dedicato.",
    ))

    r = 5
    for label, val in rows:
        ws.cell(row=r, column=2, value=label).font = bold_font
        c = ws.cell(row=r, column=3, value=val)
        c.font = normal_font
        c.alignment = wrap
        r += 1
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 90

    # ---- Fogli dati ----
    def make_sheet(name, lots):
        s = wb.create_sheet(name)
        s.sheet_view.showGridLines = False
        for i, h in enumerate(COLUMNS, start=1):
            c = s.cell(row=1, column=i, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            c.border = border
        s.freeze_panes = "A2"
        s.column_dimensions["A"].hidden = True  # colonna ID tecnica, nascosta ma leggibile via API
        for ridx, l in enumerate(lots, start=2):
            values = [
                l["id"], l["title"], l["city"], l["province"], f"{l['tribunale']} — {l['procedura']}",
                l["n_deserte"], l["prima_asta"], l["prossima_asta"],
                eur(l["prezzo_prima_asta"]), eur(l["prezzo_corrente"]), eur(l["min_offer"]),
                eur(l.get("cauzione")), "", None, None,
            ]
            for cidx, val in enumerate(values[:13], start=1):
                cell = s.cell(row=ridx, column=cidx, value=val)
                cell.font = normal_font
                cell.alignment = wrap
                cell.border = border
            link_cell = s.cell(row=ridx, column=14, value="Scheda annuncio")
            link_cell.hyperlink = l.get("page_url")
            link_cell.font = link_font
            link_cell.border = border
            docs = l.get("doc_links") or []
            docs_cell = s.cell(row=ridx, column=15, value="; ".join(docs) if docs else "Nessun documento trovato")
            docs_cell.font = normal_font
            docs_cell.alignment = wrap
            docs_cell.border = border
        widths = [8, 30, 14, 12, 32, 10, 12, 12, 16, 16, 14, 12, 26, 20, 34]
        for i, w in enumerate(widths, start=1):
            s.column_dimensions[get_column_letter(i)].width = w
        return s

    veicoli = [l for l in eligible if "AUTOVEICOLI" in (l["category"] or "")]
    arte = [l for l in eligible if l not in veicoli]
    make_sheet("Veicoli", veicoli)
    make_sheet("Arte-Oreficeria-Antiquariato", arte)
    return wb


# --------------------------------------------------------------------------------------
# STEP 7 - Google Drive: trova il report precedente + carica quello nuovo
# --------------------------------------------------------------------------------------

def get_drive_service():
    from googleapiclient.discovery import build
    from google.oauth2 import service_account

    cred_path = os.environ.get("GDRIVE_CREDENTIALS_FILE")
    if not cred_path or not os.path.exists(cred_path):
        return None
    scopes = ["https://www.googleapis.com/auth/drive"]
    creds = service_account.Credentials.from_service_account_file(cred_path, scopes=scopes)
    return build("drive", "v3", credentials=creds)


def fetch_previous_report_bytes(drive, folder_id):
    if drive is None or not folder_id:
        return None
    query = (
        f"'{folder_id}' in parents and name contains '{REPORT_PREFIX}' "
        "and trashed = false"
    )
    resp = drive.files().list(
        q=query, orderBy="createdTime desc", pageSize=5,
        fields="files(id, name, createdTime)",
    ).execute()
    files = resp.get("files", [])
    if not files:
        return None
    file_id = files[0]["id"]
    request = drive.files().get_media(fileId=file_id)
    buf = io.BytesIO()
    from googleapiclient.http import MediaIoBaseDownload
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buf.seek(0)
    return buf.read()


def upload_report(drive, folder_id, local_path):
    if drive is None or not folder_id:
        print("Credenziali Google Drive non configurate: salto l'upload, il file resta solo in ./output/")
        return
    from googleapiclient.http import MediaFileUpload
    file_metadata = {"name": os.path.basename(local_path), "parents": [folder_id]}
    media = MediaFileUpload(
        local_path,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    drive.files().create(body=file_metadata, media_body=media, fields="id").execute()
    print(f"Caricato su Google Drive: {os.path.basename(local_path)}")


# --------------------------------------------------------------------------------------
# MAIN
# --------------------------------------------------------------------------------------

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print("1/5 - Cattura token dal sito (browser headless)...")
    typesense_key, asta_token = get_live_tokens()

    print("2/5 - Ricerca lotti su Typesense...")
    raw_lots = search_lots(typesense_key)
    print(f"     trovati {len(raw_lots)} lotti (veicoli + arte/oreficeria, MO+BO+RE, in vendita)")

    print("3/5 - Recupero storico di ogni lotto e calcolo idoneita'...")
    eligible = build_eligible_lots(raw_lots, asta_token)
    print(f"     {len(eligible)} lotti idonei (>= {MIN_ASTE_DESERTE} aste deserte)")

    print("4/5 - Recupero dettagli pagina (cauzione, documenti)...")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        for lot in eligible:
            domain = domain_for(lot["tribunale"])
            permalink = lot["permalink"] or ""
            url = f"{domain}/{permalink}" if permalink else None
            lot["page_url"] = url
            if url:
                details = scrape_lot_page(page, url)
                lot["cauzione"] = details.get("cauzione")
                lot["doc_links"] = details.get("links", [])
            else:
                lot["cauzione"] = None
                lot["doc_links"] = []
        browser.close()

    print("5/5 - Confronto con report precedente e caricamento su Drive...")
    drive = get_drive_service()
    folder_id = os.environ.get("GDRIVE_FOLDER_ID")
    previous_bytes = fetch_previous_report_bytes(drive, folder_id)
    previous_ids = extract_ids_from_previous_report(previous_bytes) if previous_bytes else set()
    current_ids = {str(l["id"]) for l in eligible}
    new_ids = current_ids - previous_ids if previous_bytes else set()
    removed_ids = previous_ids - current_ids if previous_bytes else set()

    wb = build_workbook(eligible, new_ids, removed_ids, previous_found=bool(previous_bytes))
    filename = f"{REPORT_PREFIX}_{date.today().isoformat()}.xlsx"
    local_path = os.path.join(OUTPUT_DIR, filename)
    wb.save(local_path)
    print(f"File salvato in locale: {local_path}")

    upload_report(drive, folder_id, local_path)

    print("Fatto.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"ERRORE: {e}", file=sys.stderr)
        raise
